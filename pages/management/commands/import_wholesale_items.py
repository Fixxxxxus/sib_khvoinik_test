"""Импорт позиций оптового каталога /opt/ из прайса заказчика (xlsx).

Файл заказчика устроен так: слева таблица позиций (№, наименование, кол-во,
единица, розница, категория), справа отдельным блоком уровни скидок. Скидки
живут в pages/wholesale_pricing.py и сюда не попадают - команда читает только
левую таблицу.

    python manage.py import_wholesale_items /path/top20.xlsx
    python manage.py import_wholesale_items /path/top20.xlsx --dry-run
    python manage.py import_wholesale_items /path/top20.xlsx --drop-demo

Разбор наименования: заказчик пишет размер прямо в названии («Дерен белый
h 40-60 С2/3», «Липа мелколистная Tilia cordata h 1,8-2,5 (ком+сетка, D500)»).
Команда режет строку по первому размерному маркеру: слева остаётся ботаническое
имя, справа - размер. Сезонные хвосты («Осень 2026», «Н/С») уходят в описание.

Повторный запуск обновляет позиции, а не плодит дубли: слаг транслитерируется
из очищенного названия и стабилен между запусками.
"""

from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError

from pages.models import WholesaleItem, WholesaleSection
from pages.utils_slug import ascii_slugify

# Канонические разделы: регистр и пробелы в файле гуляют («Хвойные », «ДЕРЕВЬЯ»).
SECTIONS = {
    "деревья": {"title": "Деревья", "slug": "derevya", "sort_order": 10},
    "кустарники": {"title": "Кустарники", "slug": "kustarniki", "sort_order": 20},
    "хвойные": {"title": "Хвойные", "slug": "hvoynye", "sort_order": 30},
}

# Размерные маркеры в наименовании. Первый сработавший и есть начало размера.
SIZE_PATTERNS = (
    # высота: «h 40-60», «h 2,0-2,5», «h 1.8-2.5»
    re.compile(r"\bh\s*\d", re.IGNORECASE),
    # ком с сеткой: «(ком+сетка, D500)», «/ ком-сетка»
    re.compile(r"[(/]?\s*ком\s*[+\-]\s*сетка", re.IGNORECASE),
    # контейнер: «С2/3», «C20», «С5/7,5» (буква и кириллическая, и латинская)
    re.compile(r"(?<![A-Za-zА-Яа-я])[СC]\s?\d", re.IGNORECASE),
)

# Сезон отгрузки и пометка «Н/С» - это не размер, им место в описании.
SEASON_RE = re.compile(r"\b(осень|весна|лето|зима)\s*20\d{2}", re.IGNORECASE)
NS_RE = re.compile(r"(?<!\w)Н\s*/\s*С(?!\w)", re.IGNORECASE)

ASSORTMENT_RE = re.compile(r"в\s+ассортименте", re.IGNORECASE)

STOCK_MANY = {"много"}


def _clean(text) -> str:
    """Пробелы, неразрывные пробелы и мусорные восклицательные знаки."""
    value = "" if text is None else str(text)
    value = value.replace("\xa0", " ").replace(" ", " ")
    value = value.replace("!", " ")
    return re.sub(r"\s+", " ", value).strip()


def parse_price(raw) -> Decimal:
    """Цена и из «11 900,00», и из числа. Не разобрали - поднимаем ValueError."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        raise ValueError("цена пустая")
    if isinstance(raw, (int, float, Decimal)):
        value = Decimal(str(raw))
    else:
        text = str(raw).replace("\xa0", "").replace(" ", "").replace("₽", "")
        text = text.replace("руб.", "").replace("руб", "").replace(",", ".")
        try:
            value = Decimal(text)
        except InvalidOperation as exc:
            raise ValueError(f"цену «{raw}» не разобрать") from exc
    if value <= 0:
        raise ValueError(f"цена «{raw}» не больше нуля")
    return value.quantize(Decimal("0.01"))


def parse_availability(raw, unit: str) -> str:
    """Количество из файла: число, «много» или пусто.

    «много» и «МНОГО» - это не число, а текстовое наличие. Пустая ячейка
    (так у хвойных) значит «уточняйте», а не «нет в наличии».
    """
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return "уточняйте"
    if isinstance(raw, (int, float, Decimal)):
        count = int(Decimal(str(raw)))
        if count <= 0:
            return "уточняйте"
        return f"в наличии {count} {unit or 'шт'}"
    text = _clean(raw)
    if text.lower() in STOCK_MANY:
        return "в наличии"
    digits = re.fullmatch(r"\d+", text.replace(" ", ""))
    if digits:
        return f"в наличии {int(digits.group())} {unit or 'шт'}"
    return text


def split_title_and_size(raw_name: str) -> tuple[str, str, list[str]]:
    """Ботаническое имя, размер и заметки (сезон, «Н/С») из одного наименования."""
    name = _clean(raw_name)
    if not name:
        raise ValueError("пустое наименование")

    notes: list[str] = []
    season = SEASON_RE.search(name)
    if season:
        notes.append(f"Отгрузка: {season.group(0).lower()}.")
        name = _clean(SEASON_RE.sub(" ", name))
    if NS_RE.search(name):
        notes.append("Отметка заказчика: Н/С.")
        name = _clean(NS_RE.sub(" ", name))

    cut = None
    for pattern in SIZE_PATTERNS:
        match = pattern.search(name)
        if match and (cut is None or match.start() < cut):
            cut = match.start()

    if cut is None:
        return _normalize_title(name), "", notes

    title = _normalize_title(name[:cut])
    size = _clean(name[cut:]).strip(" ,/;-")
    size = re.sub(r"\(\s+", "(", re.sub(r"\s+\)", ")", size))
    if not title:
        # Размерный маркер съел всё название: лучше отдать название целиком,
        # чем сохранить позицию без имени.
        return _normalize_title(name), "", notes
    return title, size, notes


def _normalize_title(text: str) -> str:
    title = _clean(text).strip(" ,/;-")
    title = ASSORTMENT_RE.sub("в ассортименте", title)
    return title


def section_for(raw_category: str) -> dict:
    """Раздел по колонке категории: регистр и пробелы нормализуем."""
    key = _clean(raw_category).lower().replace("ё", "е")
    if not key:
        raise ValueError("пустая категория")
    if key in SECTIONS:
        return SECTIONS[key]
    title = key.capitalize()
    return {"title": title, "slug": ascii_slugify(title) or "razdel", "sort_order": 100}


def _header_map(sheet) -> dict:
    """Колонки по заголовкам: у единицы измерения заголовка нет, берём соседнюю."""
    for row in sheet.iter_rows(min_row=1, max_row=10):
        titles = {_clean(cell.value).lower(): cell.column for cell in row if cell.value}
        if "наименование" in titles:
            columns = {
                "name": titles.get("наименование"),
                "quantity": titles.get("кол-во") or titles.get("количество"),
                "price": titles.get("розница") or titles.get("цена"),
                "category": titles.get("категория"),
            }
            if columns["quantity"]:
                columns["unit"] = columns["quantity"] + 1
            return {"row": row[0].row, **columns}
    raise CommandError("В листе не найдена шапка таблицы позиций (колонка «Наименование»).")


class Command(BaseCommand):
    help = "Импорт позиций оптового каталога /opt/ из прайса заказчика (xlsx)."

    def add_arguments(self, parser):
        parser.add_argument("path", help="Путь к файлу xlsx с прайсом заказчика.")
        parser.add_argument(
            "--sheet",
            default=None,
            help="Имя листа. По умолчанию первый лист книги.",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Сухой прогон: разобрать файл и показать результат, в БД не писать.",
        )
        parser.add_argument(
            "--drop-demo",
            action="store_true",
            help="Заодно удалить демо-позиции каркаса (is_demo=True).",
        )

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - зависимость есть в requirements
            raise CommandError("Нужен openpyxl: pip install openpyxl") from exc

        path = options["path"]
        dry = options["dry_run"]
        try:
            book = openpyxl.load_workbook(path, data_only=True)
        except Exception as exc:
            raise CommandError(f"Файл не открывается: {exc}") from exc

        sheet = book[options["sheet"]] if options["sheet"] else book.worksheets[0]
        columns = _header_map(sheet)

        if options["drop_demo"] and not dry:
            removed = WholesaleItem.objects.filter(is_demo=True).delete()[0]
            WholesaleSection.objects.filter(is_demo=True, items__isnull=True).delete()
            self.stdout.write(f"Демо-позиции удалены: {removed}")

        sections: dict[str, WholesaleSection] = {}
        seen_slugs: dict[tuple[str, str], str] = {}
        created = updated = skipped = 0

        for row in sheet.iter_rows(min_row=columns["row"] + 1):
            def cell(key):
                index = columns.get(key)
                return row[index - 1].value if index and index <= len(row) else None

            raw_name = cell("name")
            raw_price = cell("price")
            raw_category = cell("category")
            if raw_price is None and not _clean(raw_category):
                # Пустая строка или подзаголовок блока («Кустарники» без цены).
                continue

            line = row[0].row
            try:
                title, size, notes = split_title_and_size(raw_name)
                block = section_for(raw_category)
                price = parse_price(raw_price)
            except ValueError as exc:
                skipped += 1
                self.stderr.write(f"Строка {line}: пропущена, {exc}. Значение: «{_clean(raw_name)}»")
                continue

            unit = _clean(cell("unit")).rstrip(".") or "шт"
            availability = parse_availability(cell("quantity"), unit)
            slug = ascii_slugify(title)
            if not slug:
                skipped += 1
                self.stderr.write(f"Строка {line}: пропущена, из названия «{title}» не вышел слаг")
                continue

            key = (block["slug"], slug)
            if key in seen_slugs and seen_slugs[key] != title:
                # Два разных названия дали один слаг: разводим руками, но заметно.
                suffix = 2
                while (block["slug"], f"{slug}-{suffix}") in seen_slugs:
                    suffix += 1
                self.stderr.write(
                    f"Строка {line}: слаг «{slug}» уже занят названием «{seen_slugs[key]}», "
                    f"беру «{slug}-{suffix}»"
                )
                slug = f"{slug}-{suffix}"
                key = (block["slug"], slug)
            seen_slugs[key] = title

            description = " ".join(notes)

            if dry:
                self.stdout.write(
                    f"[сухой прогон] {block['title']} / {slug}: «{title}» | размер: {size or '-'} | "
                    f"{price} ₽ | {availability}" + (f" | {description}" if description else "")
                )
                continue

            section = sections.get(block["slug"])
            if section is None:
                section, _ = WholesaleSection.objects.get_or_create(
                    slug=block["slug"],
                    defaults={
                        "title": block["title"],
                        "sort_order": block["sort_order"],
                        "is_active": True,
                    },
                )
                # Раздел мог остаться от демо-наполнения: данные теперь реальные.
                if section.is_demo or section.title != block["title"]:
                    section.is_demo = False
                    section.title = block["title"]
                    section.save(update_fields=["is_demo", "title"])
                sections[block["slug"]] = section

            obj, is_new = WholesaleItem.objects.update_or_create(
                section=section,
                slug=slug,
                defaults={
                    "title": title,
                    "size": size,
                    "price": price,
                    "unit": unit,
                    "availability": availability,
                    "short_description": description,
                    "sort_order": line * 10,
                    "is_active": True,
                    "is_demo": False,
                },
            )
            created += int(is_new)
            updated += int(not is_new)
            self.stdout.write(f"{'создана' if is_new else 'обновлена'}: {obj.get_absolute_url()}")

        if dry:
            self.stdout.write(self.style.SUCCESS(f"Сухой прогон закончен. Пропущено строк: {skipped}."))
            return
        self.stdout.write(
            self.style.SUCCESS(
                f"Импорт закончен. Разделов: {len(sections)}, создано: {created}, "
                f"обновлено: {updated}, пропущено: {skipped}."
            )
        )
